"""Regenerate the client deck: uv run --with python-pptx==1.0.2 python formal/reports/build_presentation.py"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = Path(__file__).parent
r = Presentation()
r.slide_width, r.slide_height = Inches(13.333), Inches(7.5)
NAVY, TEAL, WHITE, BG, GRAY, GOLD = '132C40', '007D72', 'FFFFFF', 'F3F6F7', '526576', 'AE6400'
FONT = 'Liberation Sans'

def box(s, x, y, w, h, color, rounded=False):
    a = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                           Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = RGBColor.from_string(color)
    a.line.fill.background()
    return a

def text(s, x, y, w, h, value, size=22, color=NAVY, bold=False, mono=False):
    a = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f = a.text_frame; f.word_wrap = True
    f.margin_left = f.margin_right = Inches(0)
    f.margin_top = f.margin_bottom = Inches(0)
    for i, line in enumerate(value.split('\n')):
        p = f.paragraphs[0] if i == 0 else f.add_paragraph()
        p.text = line; p.font.name = 'Liberation Mono' if mono else FONT
        p.font.size = Pt(size); p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(8)
    return a

def slide(kicker, title, subtitle='', dark=False, notes=''):
    s = r.slides.add_slide(r.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(NAVY if dark else BG)
    color = WHITE if dark else NAVY
    text(s,.6,.34,11,.3,kicker.upper(),12,TEAL if not dark else '60D6C2',True)
    text(s,.6,.91,12.1,1.0,title,34,color,True)
    if subtitle: text(s,.6,1.92,12,.65,subtitle,19,'C7D4DE' if dark else GRAY)
    box(s,.6,7.04,12.1,.012,'456071' if dark else 'CCD7DD')
    text(s,.6,7.16,10,.2,'RUST → LEAN  /  DELIVERY REVIEW  /  18 SEPTEMBER 2026',9,'C7D4DE' if dark else GRAY)
    text(s,12.1,7.12,.6,.3,f'{len(r.slides):02}',12,color)
    s.notes_slide.notes_text_frame.text = notes
    return s

def card(s,x,y,w,title,body,color=TEAL):
    box(s,x,y,w,2.4,WHITE,True)
    box(s,x,y,.055,2.4,color)
    text(s,x+.23,y+.2,w-.46,.6,title,23,color,True)
    text(s,x+.23,y+.98,w-.46,1.25,body,19)

s=slide('Client delivery','From protocol code\nto checkable evidence',dark=True,
        notes='This delivery generalizes the existing Rust/Paxos pipeline. It is a local proof engineering tool, not an automatic verifier for arbitrary Rust. See ../rust/SPECIFICATIONS.md and delivery-report.md.')
text(s,.65,3.15,10.7,1.1,'One reusable Rust-to-Lean pipeline.\nReadable claims. Explicit assumptions. Replayable proofs.',27,WHITE)
for x,n,label in [( .65,'3','working examples'),(4.8,'14','example claims checked'),(8.95,'12','regression tests passed')]:
    text(s,x,5.15,3.7,.8,n,48,'60D6C2',True)
    text(s,x,6.05,3.7,.5,label,18,WHITE)

s=slide('Executive outcome','A general engine, demonstrated beyond Paxos',
        'Adding a protocol now means adding its specification and proof adapter—not changing the driver.',
        notes='Engine: ../rust_verify.py. Elaboration: ../rust_spec.py. Generic logic: ../rmverify/lean/Specification.lean. Paxos compatibility is a data-only profile. Generality means supported extraction plus expressive Lean propositions, not complete automatic proof search.')
card(s,.65,2.95,3.85,'Reusable engine','Rust library, source file or ZIP\nOne extraction and checking path')
card(s,4.75,2.95,3.85,'Readable contracts','Commented TOML or JSON\nPreview meaning and assumptions')
card(s,8.85,2.95,3.85,'Auditable results','Exact propositions and logs\nProof replay with axiom checks')
text(s,.7,6.02,11.8,.6,'Delivered: code, specifications, examples, documentation and an agent authoring skill.',21,TEAL,True)

s=slide('Authoring experience','A small spec, with the assumption in plain sight',
        'Descriptions explain intent. Structured fields determine the exact Lean proposition.',
        notes='Working file: ../rust/delivery/spec.toml. The module, predicates and proof are defined in verification/Protocol.lean. TOML does not translate English into logic. Invariants ignore temporal assumptions; eventuality is conditional on the named assumptions.')
box(s,.65,2.8,7.1,3.8,NAVY,True)
code='[system]\nmodule = "Delivery.module"\nassumptions = ["Delivery.Fair"]\n\n[[claims]]\nname = "eventual_delivery"\nkind = "eventually"\npredicate = "Delivery.Done"\nproof = "Delivery.eventually_done"'
text(s,.93,3.05,6.65,3.25,code,17,WHITE,mono=True)
text(s,8.15,2.9,4.35,.55,'1  Preview',25,TEAL,True)
text(s,8.15,3.46,4.4,.6,'Read the claim and its premises.',19)
text(s,8.15,4.16,4.35,.55,'2  Verify',25,TEAL,True)
text(s,8.15,4.72,4.4,.65,'The crate finds its spec automatically.',19)
text(s,8.15,5.45,4.35,.55,'3  Inspect & replay',25,TEAL,True)
text(s,8.15,6.01,4.4,.65,'Review the theorem, logs and evidence.',19)

s=slide('Architecture','The proof is connected to the submitted Rust',
        'A model theorem alone is not an implementation proof.',
        notes='Rust source is snapshotted. Hax 0.4.0 invokes Charon and Aeneas to generate Lean definitions. Handwritten adapters use extracted functions directly or prove refinement. The trusted compiler/extractor boundary is not itself verified. Lean 4.31.0 is used for Rust extraction compatibility.')
labels=[('01','Rust source','Owned protocol state'),('02','Extraction','Hax / Charon / Aeneas'),('03','Lean code','Generated definitions'),('04','Proof adapter','Semantics + contracts'),('05','Kernel check','Theorem + axiom audit')]
for i,(n,title,body) in enumerate(labels):
    x=.65+i*2.48
    box(s,x,3.2,2.24,2.1,WHITE,True)
    text(s,x+.16,3.38,1.9,.4,n,17,TEAL,True)
    text(s,x+.16,3.98,1.95,.65,title,20,NAVY,True)
    text(s,x+.16,4.65,1.95,.5,body,14,GRAY)
    if i<4: text(s,x+2.29,4.0,.2,.4,'›',23,TEAL,True)
text(s,.7,5.8,11.7,.85,'Saved evidence: source snapshot + specification + generated Lean + proofs + logs + hashes.',23,TEAL,True)

s=slide('Coverage','Three examples exercise different kinds of proof',
        'All use the same driver. The example claims are checked without a finite execution-length bound.',
        notes='Paxos has four claims, delivery five and Pedersen five. Delivery includes a theorem disproving unconditional eventuality; the separate negative pipeline test returns refuted for that unconditional claim. Pedersen uses exact finite random-tape probabilities, not asymptotic complexity.')
rows=[('Example','Checked claims','What this demonstrates'),('Paxos','4 / 4','Agreement, local progress, eventual choice, refinement'),('Delivery','5 / 5','All temporal forms; why fairness is needed'),('Pedersen','5 / 5','Correctness, hiding, binding reduction and bound')]
for i,row in enumerate(rows):
    y=2.95+i*.75
    box(s,.65,y,12,.69,NAVY if i==0 else WHITE)
    for x,w,val in zip([.88,3.42,5.65],[2.3,2.1,6.55],row):
        text(s,x,y+.15,w,.43,val,17 if i else 16,WHITE if i==0 else NAVY,i==0)
text(s,.7,6.27,11.6,.5,'“Proved” applies to each precise proposition and its stated premises.',22,TEAL,True)

s=slide('Paxos result','Paxos: agreement and conditional progress',
        'Verified scope: the three-acceptor, single-decree Rust protocol core.',
        notes='See ../rust/paxos/spec.json and verification/Protocol.lean. Safety: successful learner certificates agree over all reachable states. Liveness: eventual quorum choice, with enabled local handlers, under a stable responsive quorum and weak fairness. Authentic retained messages, globally unique ballots and serialized state are part of the model/runtime interface. No Byzantine nodes.')
card(s,.65,2.9,5.88,'Safety','Successful learner certificates agree.\nNo fairness premise weakens this invariant.')
card(s,6.77,2.9,5.88,'Conditional liveness','A stable, fair quorum eventually chooses.\nEnabled local computations also succeed.')
box(s,.65,5.72,12,.88,'E8EEF1',True)
text(s,.9,5.9,11.5,.62,'Outside the proof: sockets, recovery, leader election and eventual client receipt.',21,GRAY)

s=slide('Pedersen result','Pedersen: perfect hiding and a binding reduction',
        'The security theorems are linked to extracted generic Rust commitment and verification functions.',
        notes='See ../rust/pedersen/README.md and verification/{Model,Security}.lean. Perfect hiding assumes a finite scalar field, uniform fresh randomness and a generator H. Binding derives a discrete logarithm from distinct accepted openings and transfers an explicit bound on that reduction’s success probability. BackendLaws are theorem premises. No concrete curve or RNG is verified. Background: https://arxiv.org/abs/1705.05897; its EasyCrypt proofs are not imported.')
box(s,.65,2.75,12,.77,NAVY,True)
text(s,.95,2.91,11.45,.48,'C = m · G + r · H',27,WHITE,True)
card(s,.65,3.87,5.88,'Hiding','Any two messages give identical commitment distributions under uniform blinding.')
card(s,6.77,3.87,5.88,'Binding','Two accepted, distinct openings yield log_G(H); a hardness bound transfers to forgery.')
text(s,.7,6.55,11.9,.3,'Scope: algebraic scheme + backend contracts; no curve, RNG, constant-time or asymptotic proof.',15,GOLD,True)

s=slide('Validation','Proof failures are tested, too',
        '12 regression tests passed in 450.6 seconds on this workstation.',
        notes='Command: RMVERIFY_RUST_INTEGRATION=1 python3 -m unittest formal.test_rust_pipeline -v. Run completed 2026-09-18. Includes proof replay, ZIP and standalone Rust inputs, TOML/JSON validation, unsafe archive rejection and source mutations. Runtime is an observed local validation duration, not a performance SLA.')
rows=[('Check','Observed result'),('Incorrect Paxos highest-ballot selector','Not proved'),('Paxos proposer always rejects','Unknown; liveness proof breaks'),('Pedersen blinding replaced with message','Unknown; refinement proof breaks'),('Delivery without fairness','Refuted, with a checked negation'),('A proof containing sorry','Rejected by axiom audit')]
for i,(a,b) in enumerate(rows):
    y=2.75+i*.58
    box(s,.65,y,12,.53,NAVY if i==0 else WHITE)
    text(s,.9,y+.09,6.45,.4,a,17,WHITE if i==0 else NAVY,i==0)
    text(s,7.4,y+.09,5,.4,b,17,WHITE if i==0 else TEAL,i==0)
text(s,.7,6.43,11.8,.4,'Rust unit tests passed; shared specification semantics compiled on Lean 4.31 and 4.32.',17,GRAY)

s=slide('Result semantics','A failed proof is not a counterexample',
        'The engine reports evidence it actually established.',
        notes='The report records exact Lean statements, theorem names, assumptions, diagnostics and log paths. Standard allowed axioms are propext, Classical.choice and Quot.sound. Unsupported extraction and input/tool errors are separate statuses. Replay checks hashes and proofs but does not re-extract Rust or verify the compiler.')
card(s,.65,2.95,3.85,'PROVED','Lean checked the proposition and its transitive axioms.')
card(s,4.75,2.95,3.85,'REFUTED','Lean checked the negation of that exact proposition.')
card(s,8.85,2.95,3.85,'UNKNOWN','Neither proof was established, or a proof/build timed out.',GOLD)
text(s,.7,5.98,11.9,.7,'Extraction failures and input/tooling errors are reported separately.',22,GRAY)

s=slide('Usability and modularity','A clear division of work for people and agents',
        'The interface removes repetitive claim syntax while keeping proof obligations visible.',
        notes='Read ../rust/SPECIFICATIONS.md and ../skills/write-verification-spec/SKILL.md. The skill teaches invariants, conditional liveness, source refinement, vacuity checks, proof audit, mutation tests and reporting. It is repository-local and does not alter global agent configuration.')
card(s,.65,2.95,3.85,'Protocol author','Write Rust and observable predicates.\nChoose the property kind.')
card(s,4.75,2.95,3.85,'Proof author / agent','State environment assumptions.\nSupply the adapter and lemmas.')
card(s,8.85,2.95,3.85,'Engine','Extract, elaborate, check and audit.\nSave evidence for replay.')
text(s,.7,6.05,11.7,.7,'General Lean claims cover richer mathematics; automatic proof search remains deliberately limited.',21,TEAL,True)

s=slide('Next deployment decisions','Next: connect the verified core to a real runtime',
        'Recommended priorities for turning the verified core into a deployed system.',
        notes='These are proposed follow-up scopes, not completed deliverables. General protocol composition and game-theory automation are not claimed. A hosted verifier also needs process isolation because Cargo build scripts and Lean metaprograms can execute code.')
for y,n,title,body in [(2.85,'01','Connect one real runtime','Prove the transport, ownership and recovery interface matches the model.'),(4.02,'02','Instantiate the cryptographic backend','Discharge group-operation contracts; verify parameter handling and randomness.'),(5.19,'03','Operationalize replay in CI','Pin toolchains, cache dependencies and retain proof evidence per revision.')]:
    text(s,.7,y,.7,.6,n,29,TEAL,True)
    text(s,1.6,y,10.8,.5,title,24,NAVY,True)
    text(s,1.6,y+.55,10.8,.6,body,19,GRAY)

s=slide('Handover','Start with the delivery example',dark=True,
        notes='Quickstart: ../rust/SPECIFICATIONS.md. Agent skill: ../skills/write-verification-spec/SKILL.md. Detailed report and validation commands: delivery-report.md. Run commands from repository root. For installation use python3 formal/rust_verify.py --install. Source repository: https://github.com/Flecart/paxos/tree/formal/counter-lean .')
text(s,.7,2.55,11.9,.5,'Preview → verify → replay',28,'60D6C2',True)
text(s,.7,3.42,11.9,1.8,'python3 formal/rust_verify.py --explain-spec \\\n  formal/rust/delivery/spec.toml\npython3 formal/rust_verify.py formal/rust/delivery\npython3 formal/rust_verify.py --recheck <evidence-directory>',19,WHITE,mono=True)
text(s,.7,5.6,11.8,.9,'Read: formal/rust/SPECIFICATIONS.md\nAgent guide: formal/skills/write-verification-spec/SKILL.md',19,'C7D4DE')
r.core_properties.title = 'Rust to Lean — Verification Pipeline Delivery'
r.core_properties.subject = 'Client delivery review: interface, proof results, validation and scope'
r.core_properties.author = 'Verification engineering'
r.save(OUT / 'rust-verification-delivery.pptx')
print(f'Wrote {len(r.slides)} slides to {OUT / "rust-verification-delivery.pptx"}')
